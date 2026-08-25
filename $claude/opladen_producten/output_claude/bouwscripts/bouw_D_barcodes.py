"""Bouwt import_D_barcodes.xlsx: de EAN op elke variant.

Draaien vanuit output_claude/, nadat de hoofdproducten geladen zijn en de
negende export gedraaid is. Zie opladen_karntner_s27.html, stap D.

Nodig:
  koppeltabel_D_barcodes.xlsx   interne referentie -> EAN, uit het bronbestand
  karntner_varianten.xls        de negende export (externe ID + referentie + barcode)
  karntner_exports.xls          tabblad 3, om te toetsen of een EAN al bezet is

De sleutel tussen de twee werelden is de interne referentie (175-10-36):
artikelnummer, kleurnummer, maat. Die is per databank gelijk, de externe ID's
niet -- die worden pas bij de export aangemaakt.
"""

import collections
import xml.etree.ElementTree as ET

import openpyxl

NS = '{urn:schemas-microsoft-com:office:spreadsheet}'


def blad(pad, naam=None):
    """Leest een tabblad uit een Excel-2003-XML-bestand (wat de serveractie maakt).

    Geeft (kop, rijen) terug, met elke rij op de breedte van de kop gebracht --
    export_data() laat lege cellen achteraan gewoon weg.
    """
    root = ET.parse(pad).getroot()
    for ws in root.findall(NS + 'Worksheet'):
        if naam is None or ws.get(NS + 'Name') == naam:
            rows = [[(c.find(NS + 'Data').text or '') if c.find(NS + 'Data') is not None else ''
                     for c in row.findall(NS + 'Cell')]
                    for row in ws.findall('.//' + NS + 'Row')]
            breedte = len(rows[0])
            return rows[0], [r + [''] * (breedte - len(r)) for r in rows[1:]]
    raise SystemExit('tabblad %r niet gevonden in %s' % (naam, pad))


wb = openpyxl.load_workbook('koppeltabel_D_barcodes.xlsx')
koppeltabel = {str(r[0]).strip(): str(r[1]).strip()
               for r in wb['koppeltabel'].iter_rows(min_row=2, values_only=True) if r[0]}

_, varianten = blad('karntner_varianten.xls')

# alle barcodes die al bij een Karntner-variant in productie staan, met de
# referentie erbij -- zo is een botsing met zichzelf te onderscheiden van een
# botsing met een ander product
kop3, bestaande = blad('karntner_exports.xls', '3 varianten')
i_bc, i_dc = kop3.index('barcode'), kop3.index('default_code')
bezet = {r[i_bc]: r[i_dc] for r in bestaande if r[i_bc]}

rijen, overgeslagen, botsing = [], [], []
for ext_id, referentie, barcode_nu in varianten:
    ean = koppeltabel.get(referentie)
    if not ean:
        overgeslagen.append((referentie, 'geen barcode aangeleverd'))
    elif barcode_nu == ean:
        overgeslagen.append((referentie, 'staat er al op'))
    elif ean in bezet and bezet[ean] != referentie:
        botsing.append((referentie, ean, bezet[ean]))
    else:
        rijen.append((ext_id, ean))

if botsing:
    raise SystemExit('EAN al in gebruik bij een ander product:\n' +
                     '\n'.join('  %s -> %s, bezet door %s' % b for b in botsing))

nw = openpyxl.Workbook()
ws = nw.active
ws.title = 'import'
ws.append(['id', 'barcode'])
for rij in rijen:
    ws.append(list(rij))
ws.column_dimensions['A'].width = 46
ws.column_dimensions['B'].width = 18
for cel in ws['B'][1:]:
    cel.number_format = '@'  # anders maakt Excel er een getal van

wsr = nw.create_sheet('niet mee')
wsr.append(['interne referentie', 'reden'])
for rij in sorted(overgeslagen):
    wsr.append(list(rij))
wsr.column_dimensions['A'].width = 22
wsr.column_dimensions['B'].width = 28

nw.save('import_D_barcodes.xlsx')
print('import_D_barcodes.xlsx: %d rijen' % len(rijen))
print('niet mee:', dict(collections.Counter(r[1] for r in overgeslagen)))
